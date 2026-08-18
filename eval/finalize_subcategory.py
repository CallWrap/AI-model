"""소분류 확정 골든셋 생성 — 검수 엑셀과 제외 규칙을 반영해 학습용 정답을 만든다.

  uv run --with openpyxl python eval/finalize_subcategory.py

입력
  eval/datasets/golden_category.jsonl        대본 + 중분류 정답 (사람 검수 완료)
  eval/datasets/gold_draft.jsonl             소분류 라벨 초안
  data/scenarios/subcategory_review.xlsx     계층 위반 15건의 사람 검수 결과

확정 규칙
  1. 검수 엑셀에서 '수정할 소분류'를 채웠으면 그 값 (source=human)
  2. 계층이 이미 맞는 건은 초안값 유지 (source=draft) — 검수 전이므로 완전하지 않다
  3. 검수 대상인데 빈칸이면 "맞는 소분류 없음"으로 보고 제외 (2026-08-18 결정)
  4. 스코프 밖 사업(청소년산모·발달재활)은 제외 — 과제는 장애인활동지원 한정

제외한 건도 중분류 골든셋에는 그대로 남는다. 소분류 학습에서만 빠진다.
"""

import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from callwrap_ai.schemas import SUBCATEGORIES_BY_CATEGORY

AI = Path(__file__).resolve().parents[1]
DATASETS = AI / "eval/datasets"
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else AI / "data/scenarios/subcategory_review.xlsx"
OUT = DATASETS / "golden_subcategory.jsonl"  # 커밋됨 (명세 4.2)

BY_CAT = {c.value: {s.value for s in subs} for c, subs in SUBCATEGORIES_BY_CATEGORY.items()}

# 장애인활동지원이 아닌 다른 바우처 사업. 우리 분류 체계로 담을 수 없다.
OUT_OF_SCOPE = {"scn-115": "청소년산모", "scn-128": "청소년산모",
                "scn-019": "청소년산모", "scn-121": "발달재활"}


def read_review() -> dict[str, str | None]:
    """검수 엑셀 -> {id: 수정할 소분류}. 빈칸은 None (= 맞는 소분류 없음)."""
    if not XLSX.exists():
        sys.exit(f"검수 엑셀이 없습니다: {XLSX}\n"
                 "  uv run --with openpyxl python eval/review_subcategory.py 로 만든 뒤 채웁니다.")
    ws = load_workbook(XLSX).active
    return {r[0]: (r[5] or None) for r in ws.iter_rows(min_row=2, max_col=6, values_only=True) if r[0]}


def main() -> None:
    review = read_review()
    draft = {r["scenario_id"]: r for r in map(json.loads, (DATASETS / "gold_draft.jsonl").open(encoding="utf-8"))}

    rows, dropped = [], Counter()
    for line in (DATASETS / "golden_category.jsonl").open(encoding="utf-8"):
        g = json.loads(line)
        sid, cat = g["scenario_id"], g["category"]

        if sid in OUT_OF_SCOPE:
            dropped["스코프 밖"] += 1
            continue

        if sid in review:                      # 계층 위반이라 검수 대상이었던 건
            sub = review[sid]
            source = "human"
        else:
            sub = draft.get(sid, {}).get("consultation_subcategory")
            source = "draft"

        if sub is None:
            dropped["맞는 소분류 없음"] += 1
            continue
        if sub not in BY_CAT[cat]:             # 검수 후에도 어긋나면 버그다
            dropped[f"계층 위반 잔존({sid})"] += 1
            continue

        rows.append({"scenario_id": sid, "category": cat, "subcategory": sub,
                     "source": source, "text": g["text"]})

    with OUT.open("w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")

    print(f"확정 {len(rows)}건 -> {OUT.relative_to(AI)}")
    print("  라벨 출처:", dict(Counter(r["source"] for r in rows)))
    print("  제외:", dict(dropped) or "없음")
    print()
    per_cat = Counter(r["category"] for r in rows)
    print(f"{'중분류':14s} {'건수':>4s}  {'등장 소분류':>10s} / 계약")
    for c in sorted(BY_CAT, key=lambda x: -per_cat.get(x, 0)):
        seen = {r["subcategory"] for r in rows if r["category"] == c}
        print(f"{c:12s} {per_cat.get(c, 0):5d}  {len(seen):9d} / {len(BY_CAT[c])}")
    all_seen = {r["subcategory"] for r in rows}
    missing = {s for subs in BY_CAT.values() for s in subs} - all_seen
    print(f"\n전체 소분류 {len(all_seen)}/29종 등장. 0건인 {len(missing)}종:")
    for s in sorted(missing):
        print("   ", s)


if __name__ == "__main__":
    main()
