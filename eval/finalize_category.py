"""중분류 확정 골든셋 생성 — 불일치검수 엑셀에서 중분류 정답을 확정한다.

사람이 검수한 golden_review_diff.xlsx('불일치검수' 시트)를 읽어, 다음 규칙으로
중분류 정답을 확정한다:
  1. '중분류-확정' 칸을 사람이 채웠으면 그 값 (source=human)
  2. 안 채웠고 Claude==Gemini면 그 일치값 (source=ai_agree)
  3. 안 채웠는데 갈린 경우 → 미결(리포트에 표시, 골든셋 제외)

실행: uv run --with openpyxl python eval/finalize_category.py [검수엑셀경로]
      (경로 생략 시 data/scenarios/golden_review_diff.xlsx)
"""

import sys
from collections import Counter
import json
from pathlib import Path

from openpyxl import load_workbook

from callwrap_ai.schemas import ConsultationCategory

AI = Path(__file__).resolve().parents[1]
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else AI / "data/scenarios/golden_review_diff.xlsx"
OUT = AI / "data/scenarios/golden_category.jsonl"

VALID = {c.value for c in ConsultationCategory}


def main() -> None:
    ws = load_workbook(XLSX)["불일치검수"]

    # 대본 텍스트 (파인튜닝 입력)
    dlg = {}
    for name in ("scenarios_v2.jsonl", "scenarios_v1_kept.jsonl"):
        for line in (AI / "data/scenarios" / name).open(encoding="utf-8"):
            s = json.loads(line)
            cid = s.get("spec_id") or s.get("id")
            dlg[cid] = "\n".join(f"{u['speaker']}: {u['text']}" for u in s["utterances"])

    rows, undecided, invalid = [], [], []
    for r in range(2, ws.max_row + 1):
        cid = ws.cell(r, 1).value
        if not cid:
            continue
        claude = (ws.cell(r, 6).value or "").strip()
        gemini = (ws.cell(r, 7).value or "").strip()
        confirmed = (ws.cell(r, 8).value or "").strip()

        if confirmed:
            label, source = confirmed, "human"
        elif claude == gemini and claude:
            label, source = claude, "ai_agree"
        else:
            undecided.append((cid, claude, gemini))
            continue

        if label not in VALID:
            invalid.append((cid, label))
            continue
        rows.append({"scenario_id": cid, "category": label, "source": source, "text": dlg.get(cid, "")})

    with OUT.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")

    print(f"확정 {len(rows)}건 -> {OUT.name}")
    print("  출처:", dict(Counter(r["source"] for r in rows)))
    print("  중분류 분포:", dict(Counter(r["category"] for r in rows)))
    if undecided:
        print(f"\n⚠️ 미결 {len(undecided)}건 (갈렸는데 확정칸 비어있음 — 골든셋 제외):")
        for cid, c, g in undecided:
            print(f"    {cid}: Claude={c} / Gemini={g}")
    if invalid:
        print(f"\n❌ 잘못된 라벨값 {len(invalid)}건:", invalid)


if __name__ == "__main__":
    main()
